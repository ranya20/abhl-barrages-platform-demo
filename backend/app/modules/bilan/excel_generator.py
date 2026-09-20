from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from datetime import date, timedelta
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.config import BACKEND_DIR, settings
from app.modules.annonce.excel_generator import XlsxTemplateEditor, column_number
from app.modules.bilan.mappings import get_bilan_config, normalize_barrage_code
from app.modules.bilan.repository import fetch_default_bareme_points_for_barrage, register_bilan_export
from app.modules.bilan.service import build_bilan_preview, resolve_bilan_config


MONTH_NAMES_FR = {
    1: "janvier",
    2: "février",
    3: "mars",
    4: "avril",
    5: "mai",
    6: "juin",
    7: "juillet",
    8: "août",
    9: "septembre",
    10: "octobre",
    11: "novembre",
    12: "décembre",
}


def resolve_bilan_template_path(barrage_code: str) -> Path:
    code = normalize_barrage_code(barrage_code)
    config = get_bilan_config(code)
    if not config:
        raise HTTPException(status_code=404, detail=f"Barrage non configuré pour BILAN : {code}")

    template_name = config["template"]
    candidates = [
        BACKEND_DIR / "app" / "templates" / "bilan" / template_name,
        BACKEND_DIR / "templates" / "bilan" / template_name,
        BACKEND_DIR / "excel_templates" / "bilan" / template_name,
    ]

    try:
        configured = Path(settings.EXCEL_TEMPLATES_DIR)
        if not configured.is_absolute():
            configured = (BACKEND_DIR / configured).resolve()
        candidates.append(configured / "bilan" / template_name)
    except Exception:
        pass

    for path in candidates:
        if path.exists():
            return path

    raise HTTPException(
        status_code=500,
        detail=(
            f"Modèle BILAN introuvable pour {code}. Place le fichier '{template_name}' dans "
            f"{BACKEND_DIR / 'app' / 'templates' / 'bilan'}"
        ),
    )


def resolve_exports_dir() -> Path:
    try:
        path = Path(settings.DATA_EXPORTS_DIR)
    except Exception:
        path = BACKEND_DIR / "data" / "exports"

    if not path.is_absolute():
        path = (BACKEND_DIR / path).resolve()

    path.mkdir(parents=True, exist_ok=True)
    return path


def column_name(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def clear_row(
    editor: XlsxTemplateEditor,
    sheet_name: str,
    row_number: int,
    *,
    first_column: str = "A",
    last_column: str,
):
    first = column_number(first_column)
    last = column_number(last_column)
    for index in range(first, last + 1):
        editor.clear_cell(sheet_name, f"{column_name(index)}{row_number}")


def _set_optional_number(
    editor: XlsxTemplateEditor,
    sheet_name: str,
    column: str,
    row_number: int,
    value,
):
    editor.set_number(sheet_name, f"{column}{row_number}", value)


def fill_evaporation_sheet(editor: XlsxTemplateEditor, preview: dict, config: dict):
    mapping = config["evaporation"]
    sheet_name = mapping["sheet"]
    columns = mapping["columns"]
    start_row = int(mapping["start_row"])
    row_count = int(mapping["row_count"])
    month = int(preview["month"])
    year = int(preview["year"])

    editor.set_number(sheet_name, mapping["month_cell"], month)
    editor.set_number(sheet_name, mapping["year_cell"], year)

    preview_rows = {item["date_bilan"]: item for item in preview["rows"]}
    month_start = date(year, month, 1)
    next_date = date.fromisoformat(preview["next_date"])

    for offset in range(row_count):
        excel_row = start_row + offset
        day = month_start + timedelta(days=offset)
        clear_row(editor, sheet_name, excel_row, first_column="A", last_column="I")

        if day > next_date:
            continue

        editor.set_date(sheet_name, f"{columns['date']}{excel_row}", day)

        if day == next_date:
            next_cote = preview.get("next_day_cote_7h_ngm")
            _set_optional_number(editor, sheet_name, columns["cote"], excel_row, next_cote)
            if next_cote is not None:
                last_month_row = preview_rows.get(preview["date_end"])
                next_surface = None
                calculation = (last_month_row or {}).get("calculation") or {}
                computed = calculation.get("computed") or {}
                next_surface = computed.get("surface_next_km2")
                _set_optional_number(editor, sheet_name, columns["surface"], excel_row, next_surface)
            continue

        item = preview_rows.get(day.isoformat())
        if not item:
            continue

        calculation = item.get("calculation") or {}
        inputs = calculation.get("inputs") or item.get("inputs") or {}
        computed = calculation.get("computed") or item.get("computed") or {}

        _set_optional_number(editor, sheet_name, columns["cote"], excel_row, inputs.get("cote_7h_ngm"))
        _set_optional_number(editor, sheet_name, columns["surface"], excel_row, computed.get("surface_interval_km2") or computed.get("surface_km2"))
        _set_optional_number(editor, sheet_name, columns["surface_moyenne"], excel_row, computed.get("surface_moyenne_km2"))
        _set_optional_number(editor, sheet_name, columns["hauteur_bac"], excel_row, inputs.get("hauteur_bac_mm"))
        _set_optional_number(editor, sheet_name, columns["pluie"], excel_row, inputs.get("pluie_mm"))
        _set_optional_number(editor, sheet_name, columns["hauteur_evaporee"], excel_row, computed.get("hauteur_evaporee_mm"))
        _set_optional_number(editor, sheet_name, columns["hauteur_corrigee"], excel_row, computed.get("hauteur_corrigee_mm"))
        _set_optional_number(editor, sheet_name, columns["evaporation"], excel_row, computed.get("evaporation_m3"))


def fill_apport_sheet(editor: XlsxTemplateEditor, preview: dict, config: dict):
    mapping = config["apport"]
    sheet_name = mapping["sheet"]
    columns = mapping["columns"]
    start_row = int(mapping["start_row"])
    row_count = int(mapping["row_count"])
    last_column = mapping["last_column"]
    fields = config["input_fields"]
    special_columns = mapping.get("special_columns", {})
    year = int(preview["year"])
    month = int(preview["month"])
    month_start = date(year, month, 1)
    next_date = date.fromisoformat(preview["next_date"])
    preview_rows = {item["date_bilan"]: item for item in preview["rows"]}

    for offset in range(row_count):
        excel_row = start_row + offset
        day = month_start + timedelta(days=offset)
        clear_row(editor, sheet_name, excel_row, first_column="A", last_column=last_column)

        if day > next_date:
            continue

        editor.set_date(sheet_name, f"{columns['date']}{excel_row}", day)

        if day == next_date:
            next_cote = preview.get("next_day_cote_7h_ngm")
            _set_optional_number(editor, sheet_name, columns["cote"], excel_row, next_cote)
            last_month_row = preview_rows.get(preview["date_end"])
            calculation = (last_month_row or {}).get("calculation") or {}
            computed = calculation.get("computed") or {}
            _set_optional_number(editor, sheet_name, columns["volume"], excel_row, computed.get("volume_next_mm3"))
            continue

        item = preview_rows.get(day.isoformat())
        if not item:
            continue

        calculation = item.get("calculation") or {}
        inputs = calculation.get("inputs") or item.get("inputs") or {}
        computed = calculation.get("computed") or item.get("computed") or {}
        restitution_values = inputs.get("restitutions") or {}

        _set_optional_number(editor, sheet_name, columns["cote"], excel_row, inputs.get("cote_7h_ngm"))
        _set_optional_number(editor, sheet_name, columns["volume"], excel_row, computed.get("volume_interval_mm3") or computed.get("volume_mm3"))
        _set_optional_number(editor, sheet_name, columns["variation"], excel_row, computed.get("variation_reserve_mm3"))
        _set_optional_number(editor, sheet_name, columns["evaporation"], excel_row, computed.get("evaporation_m3"))

        for field in fields:
            _set_optional_number(
                editor,
                sheet_name,
                field["column"],
                excel_row,
                restitution_values.get(field["code"]),
            )

        _set_optional_number(editor, sheet_name, columns["total_restitutions"], excel_row, computed.get("total_restitutions_m3"))
        _set_optional_number(editor, sheet_name, columns["apports"], excel_row, computed.get("apports_m3"))
        _set_optional_number(editor, sheet_name, columns["debit"], excel_row, computed.get("debit_m3s"))
        _set_optional_number(editor, sheet_name, columns["pluie"], excel_row, inputs.get("pluie_mm"))

        if "IRRIGATION_CALCULEE" in special_columns:
            _set_optional_number(
                editor,
                sheet_name,
                special_columns["IRRIGATION_CALCULEE"],
                excel_row,
                computed.get("irrigation_m3"),
            )


def fill_bilan_workbook(editor: XlsxTemplateEditor, preview: dict, config: dict):
    fill_evaporation_sheet(editor, preview, config)
    fill_apport_sheet(editor, preview, config)



# ============================================================
# Export BILAN dynamique niveau 2 pour les nouveaux barrages
# ============================================================

def _clean_title(text: str | None, fallback: str) -> str:
    value = str(text or fallback).strip()
    value = value.replace("Barrage ", "").replace("barrage ", "")
    return value.upper()


def _cell_number(ws, row: int, col: int, value, number_format: str = "0.000"):
    cell = ws.cell(row=row, column=col, value=value)
    if value is not None:
        cell.number_format = number_format
    return cell


def _style_title(ws, row: int, start_col: int, end_col: int, text: str, fill, font):
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row=row, column=start_col, value=text)
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_header_cell(cell, fill, font, border):
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border


def _style_data_cell(cell, border):
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border


def _auto_width(ws):
    """
    Ajuste automatiquement la largeur des colonnes.

    Version robuste :
    - ignore les cellules fusionnées MergedCell
    - utilise l'index de colonne au lieu de column_cells[0].column_letter
    - évite l'erreur : AttributeError: 'MergedCell' object has no attribute 'column_letter'
    """
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_length = 0

        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)

            # Les cellules issues d'une fusion n'ont pas toujours les mêmes attributs.
            if cell.__class__.__name__ == "MergedCell":
                continue

            value = cell.value
            if value is None:
                continue

            text = str(value)
            for part in text.split("\n"):
                max_length = max(max_length, len(part))

        if max_length <= 0:
            width = 10
        else:
            width = min(max(max_length + 2, 10), 35)

        ws.column_dimensions[col_letter].width = width


def _row_lookup(preview: dict) -> dict[str, dict]:
    return {row["date_bilan"]: row for row in preview.get("rows", [])}


def _computed_for_row(row: dict) -> dict:
    calculation = row.get("calculation") or {}
    return calculation.get("computed") or row.get("computed") or {}


def _inputs_for_row(row: dict) -> dict:
    calculation = row.get("calculation") or {}
    return calculation.get("inputs") or row.get("inputs") or {}


def _status_for_row(row: dict) -> str:
    calculation = row.get("calculation") or {}
    return calculation.get("status") or row.get("statut") or ""


def _create_apport_sheet(wb, preview: dict, config: dict, styles: dict):
    ws = wb.active
    ws.title = "APPORT"

    fields = config["input_fields"]
    barrage = preview["barrage"]
    title = _clean_title(barrage.get("barrage_nom"), preview["barrage"]["barrage_code"])
    year = int(preview["year"])
    month = int(preview["month"])

    base_headers = [
        "Date",
        "Cote à 7h\n(NGM)",
        "Cote suivante\n(NGM)",
        "Volume\n(Mm³)",
        "Volume suivant\n(Mm³)",
        "Variation\n(Mm³)",
        "Evaporation\n(m³)",
    ]
    restitution_headers = [field["label"] for field in fields]
    tail_headers = ["Total restitutions\n(m³)", "Apports\n(m³)", "Débit\n(m³/s)", "Pluie\n(mm)", "Statut"]
    headers = base_headers + restitution_headers + tail_headers

    end_col = len(headers)

    _style_title(ws, 1, 1, end_col, f"BILAN MENSUEL - {title}", styles["blue"], styles["title_font"])
    _style_title(ws, 2, 1, end_col, f"{MONTH_NAMES_FR[month].upper()} {year}", styles["blue"], styles["subtitle_font"])

    header_row = 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        _style_header_cell(cell, styles["blue"], styles["header_font"], styles["border"])

    start_row = header_row + 1
    for index, row in enumerate(preview.get("rows", []), start=start_row):
        inputs = _inputs_for_row(row)
        computed = _computed_for_row(row)
        restitutions = inputs.get("restitutions") or {}

        values = [
            row.get("date_bilan"),
            inputs.get("cote_7h_ngm"),
            inputs.get("cote_suivante_ngm"),
            computed.get("volume_interval_mm3") or computed.get("volume_mm3"),
            computed.get("volume_next_mm3"),
            computed.get("variation_reserve_mm3"),
            computed.get("evaporation_m3"),
        ]
        for field in fields:
            values.append(restitutions.get(field["code"], 0))
        values.extend(
            [
                computed.get("total_restitutions_m3"),
                computed.get("apports_m3"),
                computed.get("debit_m3s"),
                inputs.get("pluie_mm"),
                _status_for_row(row),
            ]
        )

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            _style_data_cell(cell, styles["border"])
            if col != 1 and isinstance(value, (int, float)):
                cell.number_format = "0.000"

    total_row = start_row + len(preview.get("rows", [])) + 1
    ws.cell(row=total_row, column=1, value="Total").font = styles["bold_font"]

    for col in range(2, end_col + 1):
        col_letter = get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f"=SUM({col_letter}{start_row}:{col_letter}{total_row-2})")
        ws.cell(row=total_row, column=col).font = styles["bold_font"]
        ws.cell(row=total_row, column=col).border = styles["border"]
        ws.cell(row=total_row, column=col).number_format = "0.000"

    ws.freeze_panes = "A5"
    _auto_width(ws)


def _create_evaporation_sheet(wb, preview: dict, styles: dict):
    ws = wb.create_sheet("EVAPORATION")
    barrage = preview["barrage"]
    title = _clean_title(barrage.get("barrage_nom"), barrage.get("barrage_code"))
    year = int(preview["year"])
    month = int(preview["month"])

    headers = [
        "Date",
        "Cote à 7h\n(NGM)",
        "Surface\n(km²)",
        "Surface moyenne\n(km²)",
        "Hauteur Bac\n(mm)",
        "Pluie\n(mm)",
        "Hauteur évaporée\n(mm)",
        "Hauteur corrigée\n(mm)",
        "Evaporation\n(m³)",
    ]

    _style_title(ws, 1, 1, len(headers), f"EVAPORATION - {title}", styles["blue"], styles["title_font"])
    _style_title(ws, 2, 1, len(headers), f"{MONTH_NAMES_FR[month].upper()} {year}", styles["blue"], styles["subtitle_font"])

    header_row = 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        _style_header_cell(cell, styles["blue"], styles["header_font"], styles["border"])

    start_row = header_row + 1
    for index, row in enumerate(preview.get("rows", []), start=start_row):
        inputs = _inputs_for_row(row)
        computed = _computed_for_row(row)
        values = [
            row.get("date_bilan"),
            inputs.get("cote_7h_ngm"),
            computed.get("surface_interval_km2") or computed.get("surface_km2"),
            computed.get("surface_moyenne_km2"),
            inputs.get("hauteur_bac_mm"),
            inputs.get("pluie_mm"),
            computed.get("hauteur_evaporee_mm"),
            computed.get("hauteur_corrigee_mm"),
            computed.get("evaporation_m3"),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            _style_data_cell(cell, styles["border"])
            if col != 1 and isinstance(value, (int, float)):
                cell.number_format = "0.000"

    ws.freeze_panes = "A5"
    _auto_width(ws)


def _create_bareme_sheet(wb, points: list[dict], preview: dict, styles: dict):
    ws = wb.create_sheet("BAREME")
    barrage = preview["barrage"]
    title = _clean_title(barrage.get("barrage_nom"), barrage.get("barrage_code"))

    headers = ["Cote NGM", "Surface km²", "Volume Mm³"]
    _style_title(ws, 1, 1, len(headers), f"BAREME - {title}", styles["blue"], styles["title_font"])

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        _style_header_cell(cell, styles["blue"], styles["header_font"], styles["border"])

    for index, point in enumerate(points, start=4):
        values = [point.get("cote_ngm"), point.get("surface_km2"), point.get("volume_mm3")]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            _style_data_cell(cell, styles["border"])
            cell.number_format = "0.000"

    _auto_width(ws)


def _create_parametres_sheet(wb, preview: dict, config: dict, styles: dict):
    ws = wb.create_sheet("PARAMETRES")
    barrage = preview["barrage"]
    rows = [
        ("Code barrage", barrage.get("barrage_code")),
        ("Nom barrage", barrage.get("barrage_nom")),
        ("Capacité normale Mm³", barrage.get("capacite_normale_mm3")),
        ("Année", preview.get("year")),
        ("Mois", preview.get("month")),
        ("Type export", "BILAN dynamique généré par la plateforme"),
        ("Nombre restitutions", len(config.get("input_fields", []))),
    ]
    ws.cell(row=1, column=1, value="PARAMETRES DU BILAN DYNAMIQUE").font = styles["title_font"]
    for index, (key, value) in enumerate(rows, start=3):
        ws.cell(row=index, column=1, value=key).font = styles["bold_font"]
        ws.cell(row=index, column=2, value=value)
    _auto_width(ws)


def generate_dynamic_bilan_workbook(db: Session, output_path: Path, preview: dict, config: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    styles = {
        "blue": PatternFill("solid", fgColor="4F81BD"),
        "title_font": Font(bold=True, color="FFFFFF", size=16),
        "subtitle_font": Font(bold=True, color="FFFFFF", size=13),
        "header_font": Font(bold=True, color="FFFFFF"),
        "bold_font": Font(bold=True),
        "border": Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        ),
    }

    _create_apport_sheet(wb, preview, config, styles)
    _create_evaporation_sheet(wb, preview, styles)

    points = fetch_default_bareme_points_for_barrage(db, preview["barrage"]["barrage_code"])
    _create_bareme_sheet(wb, points, preview, styles)
    _create_parametres_sheet(wb, preview, config, styles)

    wb.save(output_path)

def build_bilan_filename(barrage_code: str, year: int, month: int, config: dict | None = None) -> str:
    code = normalize_barrage_code(barrage_code)
    if config and config.get("dynamic"):
        template_stem = f"BILAN {code}"
    else:
        static_config = config or get_bilan_config(code)
        template_stem = Path(static_config["template"]).stem if static_config else f"BILAN {code}"
    token = uuid4().hex[:8]
    return f"{template_stem} - {year}-{month:02d} - {token}.xlsx"


def generate_bilan_excel(db: Session, barrage_code: str, year: int, month: int) -> Path:
    code, barrage, config = resolve_bilan_config(db, barrage_code)
    preview = build_bilan_preview(db, code, year, month)
    output_dir = resolve_exports_dir()
    output_path = output_dir / build_bilan_filename(code, year, month, config)

    if config.get("dynamic"):
        generate_dynamic_bilan_workbook(db, output_path, preview, config)
    else:
        template_path = resolve_bilan_template_path(code)
        editor = XlsxTemplateEditor(template_path)
        fill_bilan_workbook(editor, preview, config)
        editor.save(output_path)

    try:
        register_bilan_export(
            db,
            barrage_code=code,
            year=year,
            month=month,
            filename=output_path.name,
            output_path=output_path,
        )
        db.commit()
    except Exception:
        db.rollback()
        # L'export Excel reste valide même si la table de journalisation n'est pas disponible.

    return output_path

# === ABHL DYNAMIC BILAN V2 PATCH START ===
# Ce bloc surcharge la génération BILAN dynamique pour les nouveaux barrages.
# Il est volontairement autonome pour éviter les problèmes d'imports existants.

def _abhl_v2_to_float(value):
    if value is None:
        return None
    try:
        return float(value)
    except Exception:
        return None


def _abhl_v2_to_date(value):
    from datetime import date, datetime

    if value is None:
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except Exception:
            return value
    return value


def _abhl_v2_month_name_fr(month: int) -> str:
    names = {
        1: "JANVIER",
        2: "FÉVRIER",
        3: "MARS",
        4: "AVRIL",
        5: "MAI",
        6: "JUIN",
        7: "JUILLET",
        8: "AOÛT",
        9: "SEPTEMBRE",
        10: "OCTOBRE",
        11: "NOVEMBRE",
        12: "DÉCEMBRE",
    }
    return names.get(int(month), str(month))


def _abhl_v2_get_barrage_code(preview: dict) -> str:
    barrage = preview.get("barrage") or {}
    return barrage.get("barrage_code") or barrage.get("code") or ""


def _abhl_v2_get_barrage_title(preview: dict) -> str:
    barrage = preview.get("barrage") or {}
    name = barrage.get("barrage_nom_court") or barrage.get("barrage_nom") or barrage.get("nom") or _abhl_v2_get_barrage_code(preview)
    return str(name).strip()


def _abhl_v2_restitution_fields(preview: dict) -> list[dict]:
    fields = []

    for item in preview.get("input_fields") or []:
        role = str(item.get("role") or "").lower()
        code = item.get("code")
        if not code:
            continue

        # Tous les champs marqués restitution sont dynamiques.
        # On ignore les champs système non restitution.
        if role == "restitution" or item.get("included_in_total") is True:
            fields.append(
                {
                    "code": str(code),
                    "label": str(item.get("label") or code),
                    "unit": str(item.get("unit") or "m3"),
                    "included_in_total": bool(item.get("included_in_total", True)),
                }
            )

    return fields


def _abhl_v2_make_styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    blue = "4F81BD"
    dark_blue = "1F4E79"
    green = "92D050"
    light_blue = "D9EAF7"
    light_gray = "F2F2F2"
    yellow = "FFF2CC"
    red = "F4CCCC"
    white = "FFFFFF"
    black = "000000"

    thin = Side(style="thin", color=black)
    medium = Side(style="medium", color=black)

    return {
        "blue_fill": PatternFill("solid", fgColor=blue),
        "dark_blue_fill": PatternFill("solid", fgColor=dark_blue),
        "green_fill": PatternFill("solid", fgColor=green),
        "light_blue_fill": PatternFill("solid", fgColor=light_blue),
        "gray_fill": PatternFill("solid", fgColor=light_gray),
        "yellow_fill": PatternFill("solid", fgColor=yellow),
        "red_fill": PatternFill("solid", fgColor=red),
        "white_font": Font(color=white, bold=True, size=11),
        "title_font": Font(color=white, bold=True, size=16),
        "subtitle_font": Font(color=white, bold=True, size=13),
        "header_font": Font(color=white, bold=True, size=10),
        "bold_font": Font(bold=True),
        "normal_font": Font(size=10),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "medium_border": Border(left=medium, right=medium, top=medium, bottom=medium),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "right": Alignment(horizontal="right", vertical="center", wrap_text=True),
    }


def _abhl_v2_style_cell(cell, fill=None, font=None, border=None, alignment=None, number_format=None):
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if border is not None:
        cell.border = border
    if alignment is not None:
        cell.alignment = alignment
    if number_format is not None:
        cell.number_format = number_format


def _abhl_v2_merge_title(ws, row: int, start_col: int, end_col: int, text: str, styles, height=28):
    from openpyxl.utils import get_column_letter

    start = get_column_letter(start_col)
    end = get_column_letter(end_col)
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)

    cell = ws.cell(row=row, column=start_col)
    cell.value = text
    _abhl_v2_style_cell(
        cell,
        fill=styles["blue_fill"],
        font=styles["title_font"],
        border=styles["medium_border"],
        alignment=styles["center"],
    )

    for col in range(start_col + 1, end_col + 1):
        _abhl_v2_style_cell(
            ws.cell(row=row, column=col),
            fill=styles["blue_fill"],
            border=styles["medium_border"],
            alignment=styles["center"],
        )

    ws.row_dimensions[row].height = height


def _abhl_v2_write_header_row(ws, row: int, headers: list[str], styles):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        _abhl_v2_style_cell(
            cell,
            fill=styles["blue_fill"],
            font=styles["header_font"],
            border=styles["border"],
            alignment=styles["center"],
        )
    ws.row_dimensions[row].height = 36


def _abhl_v2_apply_table_style(ws, min_row: int, max_row: int, min_col: int, max_col: int, styles):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            _abhl_v2_style_cell(
                cell,
                border=styles["border"],
                alignment=styles["center"],
                font=styles["normal_font"],
            )


def _abhl_v2_auto_width(ws, max_width: int = 32):
    from openpyxl.utils import get_column_letter

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0

        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)

            if cell.__class__.__name__ == "MergedCell":
                continue

            value = cell.value
            if value is None:
                continue

            for part in str(value).split("\n"):
                max_len = max(max_len, len(part))

        width = min(max(max_len + 2, 10), max_width)
        ws.column_dimensions[col_letter].width = width


def _abhl_v2_freeze_and_filter(ws, freeze_cell: str, header_row: int, last_col: int, last_row: int):
    from openpyxl.utils import get_column_letter

    ws.freeze_panes = freeze_cell
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"


def _abhl_v2_status_for_row(row: dict) -> str:
    computed = row.get("computed")
    inputs = row.get("inputs") or {}

    if not computed:
        return "NON SAISI"

    required = [
        inputs.get("cote_7h_ngm"),
        inputs.get("cote_suivante_ngm"),
        inputs.get("hauteur_bac_mm"),
        inputs.get("pluie_mm"),
        computed.get("apports_m3"),
    ]

    if any(value is None for value in required):
        return "INCOMPLET"

    return "OK"


def _abhl_v2_write_guide_sheet(wb, preview: dict, styles):
    ws = wb.create_sheet("Guide d'utilisation", 0)

    title = _abhl_v2_get_barrage_title(preview)
    code = _abhl_v2_get_barrage_code(preview)
    month = int(preview.get("month") or 1)
    year = int(preview.get("year") or 1900)

    _abhl_v2_merge_title(
        ws,
        row=1,
        start_col=1,
        end_col=6,
        text=f"BILAN MENSUEL DYNAMIQUE - {title.upper()}",
        styles=styles,
        height=30,
    )

    rows = [
        ("Barrage", title),
        ("Code", code),
        ("Mois", f"{_abhl_v2_month_name_fr(month)} {year}"),
        ("Type", "BILAN mensuel généré dynamiquement pour un barrage ajouté dans la plateforme"),
        ("Source des données", "PostgreSQL - bilans_journaliers, restitutions_journalieres, bareme_points"),
        ("Important", "Les anciennes feuilles officielles restent utilisées pour les anciens barrages. Ce modèle dynamique est utilisé uniquement pour les nouveaux barrages."),
    ]

    start_row = 4
    for offset, (label, value) in enumerate(rows):
        r = start_row + offset
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=value)

        _abhl_v2_style_cell(ws.cell(row=r, column=1), fill=styles["light_blue_fill"], font=styles["bold_font"], border=styles["border"], alignment=styles["left"])
        _abhl_v2_style_cell(ws.cell(row=r, column=2), border=styles["border"], alignment=styles["left"])

        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28


def _abhl_v2_write_apport_sheet(wb, preview: dict, styles):
    """
    Feuille APPORT avec structure proche des fichiers BILAN officiels.

    Structure dynamique :
    - colonnes fixes à gauche
    - groupe "Restitutions (m³)" qui couvre les restitutions + Total
    - colonnes Apports / Débit / Pluie / Observation hors groupe restitutions

    Exemple avec 3 restitutions :
    H-J : AEPI, Fuites, EVAC
    K   : Total
    L   : Apports
    M   : Débit
    N   : Pluie
    O   : Observation

    Exemple avec 5 restitutions :
    H-L : restitutions
    M   : Total
    N   : Apports
    O   : Débit
    P   : Pluie
    Q   : Observation
    """
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("APPORT")

    title = _abhl_v2_get_barrage_title(preview)
    month = int(preview.get("month") or 1)
    year = int(preview.get("year") or 1900)
    restitution_fields = _abhl_v2_restitution_fields(preview)

    static_headers = [
        "Date",
        "Cote à 7h\n(NGM)",
        "Cote suivante\n(NGM)",
        "Volume\n(Mm³)",
        "Volume suivant\n(Mm³)",
        "Variation réserve\n(Mm³)",
        "Evaporation\n(m³)",
    ]

    static_count = len(static_headers)

    first_restitution_col = static_count + 1
    last_restitution_col = first_restitution_col + len(restitution_fields) - 1

    # Le total fait partie du groupe Restitutions, comme dans les fichiers officiels.
    total_restitution_col = first_restitution_col + len(restitution_fields)

    apports_col = total_restitution_col + 1
    debit_col = apports_col + 1
    pluie_col = debit_col + 1
    observation_col = pluie_col + 1
    last_col = observation_col

    _abhl_v2_merge_title(ws, 1, 1, last_col, f"BILAN MENSUEL - {title.upper()}", styles)
    _abhl_v2_merge_title(ws, 2, 1, last_col, f"{_abhl_v2_month_name_fr(month)} {year}", styles, height=22)

    group_row = 4
    subheader_row = 5
    data_start_row = 6

    # Style général des deux lignes d'en-tête.
    for row in [group_row, subheader_row]:
        for col in range(1, last_col + 1):
            _abhl_v2_style_cell(
                ws.cell(row=row, column=col),
                fill=styles["blue_fill"],
                font=styles["header_font"],
                border=styles["border"],
                alignment=styles["center"],
            )

    # Colonnes fixes à gauche : fusion verticale ligne 4-5.
    for col_idx, label in enumerate(static_headers, start=1):
        ws.merge_cells(
            start_row=group_row,
            start_column=col_idx,
            end_row=subheader_row,
            end_column=col_idx,
        )
        cell = ws.cell(row=group_row, column=col_idx, value=label)
        _abhl_v2_style_cell(
            cell,
            fill=styles["blue_fill"],
            font=styles["header_font"],
            border=styles["border"],
            alignment=styles["center"],
        )

    # Groupe Restitutions : il couvre toutes les restitutions + Total.
    ws.merge_cells(
        start_row=group_row,
        start_column=first_restitution_col,
        end_row=group_row,
        end_column=total_restitution_col,
    )
    group_cell = ws.cell(row=group_row, column=first_restitution_col, value="Restitutions (m³)")
    _abhl_v2_style_cell(
        group_cell,
        fill=styles["blue_fill"],
        font=styles["header_font"],
        border=styles["border"],
        alignment=styles["center"],
    )

    # Sous-colonnes restitutions dynamiques.
    for offset, field in enumerate(restitution_fields):
        col_idx = first_restitution_col + offset
        cell = ws.cell(row=subheader_row, column=col_idx, value=field["label"])
        _abhl_v2_style_cell(
            cell,
            fill=styles["blue_fill"],
            font=styles["header_font"],
            border=styles["border"],
            alignment=styles["center"],
        )

    # Sous-colonne Total dans le même groupe Restitutions.
    total_cell = ws.cell(row=subheader_row, column=total_restitution_col, value="Total")
    _abhl_v2_style_cell(
        total_cell,
        fill=styles["blue_fill"],
        font=styles["header_font"],
        border=styles["border"],
        alignment=styles["center"],
    )

    # Colonnes hors groupe Restitutions : fusion verticale ligne 4-5.
    independent_headers = [
        (apports_col, "Apports\n(m³)"),
        (debit_col, "Débit\n(m³/s)"),
        (pluie_col, "Pluie\n(mm)"),
        (observation_col, "Observation"),
    ]

    for col_idx, label in independent_headers:
        ws.merge_cells(
            start_row=group_row,
            start_column=col_idx,
            end_row=subheader_row,
            end_column=col_idx,
        )
        cell = ws.cell(row=group_row, column=col_idx, value=label)
        _abhl_v2_style_cell(
            cell,
            fill=styles["blue_fill"],
            font=styles["header_font"],
            border=styles["border"],
            alignment=styles["center"],
        )

    ws.row_dimensions[group_row].height = 28
    ws.row_dimensions[subheader_row].height = 34

    rows = preview.get("rows") or []

    for index, row in enumerate(rows):
        excel_row = data_start_row + index
        inputs = row.get("inputs") or {}
        computed = row.get("computed") or {}
        restitutions = inputs.get("restitutions") or {}

        # Colonnes fixes.
        values_by_col = {
            1: _abhl_v2_to_date(row.get("date_bilan")),
            2: _abhl_v2_to_float(inputs.get("cote_7h_ngm")),
            3: _abhl_v2_to_float(inputs.get("cote_suivante_ngm")),
            4: _abhl_v2_to_float(computed.get("volume_mm3")),
            5: _abhl_v2_to_float(computed.get("volume_jour_suivant_mm3")),
            6: _abhl_v2_to_float(computed.get("variation_reserve_mm3")),
            7: _abhl_v2_to_float(computed.get("evaporation_m3")),
        }

        # Restitutions dynamiques.
        for offset, field in enumerate(restitution_fields):
            values_by_col[first_restitution_col + offset] = _abhl_v2_to_float(
                restitutions.get(field["code"])
            ) or 0

        # Total dans le groupe Restitutions.
        values_by_col[total_restitution_col] = _abhl_v2_to_float(
            computed.get("total_restitutions_m3")
        )

        # Colonnes indépendantes.
        values_by_col[apports_col] = _abhl_v2_to_float(computed.get("apports_m3"))
        values_by_col[debit_col] = _abhl_v2_to_float(computed.get("debit_m3s"))
        values_by_col[pluie_col] = _abhl_v2_to_float(inputs.get("pluie_mm"))
        values_by_col[observation_col] = inputs.get("observation")

        for col_idx in range(1, last_col + 1):
            cell = ws.cell(row=excel_row, column=col_idx, value=values_by_col.get(col_idx))
            _abhl_v2_style_cell(cell, border=styles["border"], alignment=styles["center"])

            if col_idx == 1:
                cell.number_format = "dd/mm/yyyy"
            elif col_idx == observation_col:
                cell.alignment = styles["left"]
            else:
                cell.number_format = "#,##0.000"

        status = _abhl_v2_status_for_row(row)
        if status == "NON SAISI":
            for col_idx in range(1, last_col + 1):
                ws.cell(row=excel_row, column=col_idx).fill = styles["gray_fill"]
        elif status == "INCOMPLET":
            for col_idx in range(1, last_col + 1):
                ws.cell(row=excel_row, column=col_idx).fill = styles["yellow_fill"]

    # Ligne Total : uniquement les colonnes additives.
    total_row = data_start_row + len(rows) + 1
    ws.cell(row=total_row, column=1, value="Total")
    _abhl_v2_style_cell(
        ws.cell(row=total_row, column=1),
        fill=styles["blue_fill"],
        font=styles["white_font"],
        border=styles["border"],
        alignment=styles["center"],
    )

    additive_cols = [6, 7]  # Variation réserve, Evaporation

    if restitution_fields:
        additive_cols.extend(range(first_restitution_col, total_restitution_col + 1))
    else:
        additive_cols.append(total_restitution_col)

    additive_cols.extend([apports_col, pluie_col])

    for col_idx in range(2, last_col + 1):
        cell = ws.cell(row=total_row, column=col_idx)

        if col_idx in additive_cols:
            col_letter = get_column_letter(col_idx)
            cell.value = f"=SUM({col_letter}{data_start_row}:{col_letter}{data_start_row + len(rows) - 1})"
            cell.number_format = "#,##0.000"
        elif col_idx == debit_col:
            apports_letter = get_column_letter(apports_col)
            cell.value = f"=IFERROR({apports_letter}{total_row}/(COUNT({apports_letter}{data_start_row}:{apports_letter}{data_start_row + len(rows) - 1})*86400),\"\")"
            cell.number_format = "0.000"
        else:
            cell.value = ""

        _abhl_v2_style_cell(
            cell,
            fill=styles["light_blue_fill"],
            font=styles["bold_font"],
            border=styles["border"],
            alignment=styles["center"],
        )

    _abhl_v2_apply_table_style(ws, group_row, total_row, 1, last_col, styles)

    # Pour rester proche des modèles officiels, on fige sans imposer de filtre Excel.
    ws.freeze_panes = "A6"

    _abhl_v2_auto_width(ws)

def _abhl_v2_write_evaporation_sheet(wb, preview: dict, styles):
    ws = wb.create_sheet("EVAPORATION")

    title = _abhl_v2_get_barrage_title(preview)
    month = int(preview.get("month") or 1)
    year = int(preview.get("year") or 1900)

    headers = [
        "Date",
        "Cote à 7h\n(NGM)",
        "Surface\n(km²)",
        "Surface moyenne\n(km²)",
        "Hauteur Bac\n(mm)",
        "Pluie\n(mm)",
        "Hauteur évaporée\n(mm)",
        "Hauteur corrigée\n(mm)",
        "Evaporation\n(m³)",
    ]

    last_col = len(headers)

    _abhl_v2_merge_title(ws, 1, 1, last_col, f"EVAPORATION - {title.upper()}", styles)
    _abhl_v2_merge_title(ws, 2, 1, last_col, f"{_abhl_v2_month_name_fr(month)} {year}", styles, height=22)

    header_row = 4
    data_start_row = 5
    _abhl_v2_write_header_row(ws, header_row, headers, styles)

    rows = preview.get("rows") or []

    for index, row in enumerate(rows):
        excel_row = data_start_row + index
        inputs = row.get("inputs") or {}
        computed = row.get("computed") or {}

        values = [
            _abhl_v2_to_date(row.get("date_bilan")),
            _abhl_v2_to_float(inputs.get("cote_7h_ngm")),
            _abhl_v2_to_float(computed.get("surface_km2")),
            _abhl_v2_to_float(computed.get("surface_moyenne_km2")),
            _abhl_v2_to_float(inputs.get("hauteur_bac_mm")),
            _abhl_v2_to_float(inputs.get("pluie_mm")),
            _abhl_v2_to_float(computed.get("hauteur_evaporee_mm")),
            _abhl_v2_to_float(computed.get("hauteur_corrigee_mm")),
            _abhl_v2_to_float(computed.get("evaporation_m3")),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            _abhl_v2_style_cell(cell, border=styles["border"], alignment=styles["center"])

            if col_idx == 1:
                cell.number_format = "dd/mm/yyyy"
            else:
                cell.number_format = "#,##0.000"

        status = _abhl_v2_status_for_row(row)
        if status == "NON SAISI":
            for col_idx in range(1, last_col + 1):
                ws.cell(row=excel_row, column=col_idx).fill = styles["gray_fill"]
        elif status == "INCOMPLET":
            for col_idx in range(1, last_col + 1):
                ws.cell(row=excel_row, column=col_idx).fill = styles["yellow_fill"]

    total_row = data_start_row + len(rows) + 1
    ws.cell(row=total_row, column=1, value="Total")
    _abhl_v2_style_cell(ws.cell(row=total_row, column=1), fill=styles["blue_fill"], font=styles["white_font"], border=styles["border"], alignment=styles["center"])

    # Totaux uniquement pour les colonnes qui s'additionnent vraiment.
    additive_cols = [5, 6, 7, 8, 9]

    for col_idx in range(2, last_col + 1):
        cell = ws.cell(row=total_row, column=col_idx)

        if col_idx in additive_cols:
            from openpyxl.utils import get_column_letter

            col_letter = get_column_letter(col_idx)
            cell.value = f"=SUM({col_letter}{data_start_row}:{col_letter}{data_start_row + len(rows) - 1})"
            cell.number_format = "#,##0.000"
        else:
            cell.value = ""

        _abhl_v2_style_cell(cell, fill=styles["light_blue_fill"], font=styles["bold_font"], border=styles["border"], alignment=styles["center"])

    _abhl_v2_apply_table_style(ws, header_row, total_row, 1, last_col, styles)
    _abhl_v2_freeze_and_filter(ws, "A5", header_row, last_col, total_row)
    _abhl_v2_auto_width(ws)


def _abhl_v2_fetch_bareme_points(db, barrage_code: str) -> list[dict]:
    from sqlalchemy import text as sa_text

    rows = db.execute(
        sa_text(
            """
            SELECT
                bp.cote_ngm,
                bp.surface_km2,
                bp.volume_mm3
            FROM public.barrages b
            JOIN public.bareme_versions bv
              ON bv.barrage_id = b.id
            JOIN public.bareme_points bp
              ON bp.bareme_version_id = bv.id
            WHERE b.code = :code
              AND COALESCE(bv.actif, TRUE) = TRUE
              AND COALESCE(bv.is_default, TRUE) = TRUE
            ORDER BY bp.cote_ngm;
            """
        ),
        {"code": barrage_code},
    ).mappings().all()

    return [
        {
            "cote_ngm": _abhl_v2_to_float(row["cote_ngm"]),
            "surface_km2": _abhl_v2_to_float(row["surface_km2"]),
            "volume_mm3": _abhl_v2_to_float(row["volume_mm3"]),
        }
        for row in rows
    ]


def _abhl_v2_write_bareme_sheet(wb, db, preview: dict, styles):
    ws = wb.create_sheet("BAREME")

    title = _abhl_v2_get_barrage_title(preview)
    code = _abhl_v2_get_barrage_code(preview)
    points = _abhl_v2_fetch_bareme_points(db, code)

    headers = ["Cote\n(NGM)", "Surface\n(km²)", "Volume\n(Mm³)"]
    last_col = len(headers)

    _abhl_v2_merge_title(ws, 1, 1, last_col, f"BAREME - {title.upper()}", styles)
    _abhl_v2_write_header_row(ws, 3, headers, styles)

    for index, point in enumerate(points, start=4):
        values = [
            point.get("cote_ngm"),
            point.get("surface_km2"),
            point.get("volume_mm3"),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col_idx, value=value)
            _abhl_v2_style_cell(cell, border=styles["border"], alignment=styles["center"], number_format="#,##0.000")

    if not points:
        ws.cell(row=4, column=1, value="Aucun point de barème trouvé.")
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=3)

    _abhl_v2_auto_width(ws)


def _abhl_v2_write_controle_sheet(wb, preview: dict, styles):
    ws = wb.create_sheet("CONTROLE")

    title = _abhl_v2_get_barrage_title(preview)
    _abhl_v2_merge_title(ws, 1, 1, 5, f"CONTROLE - {title.upper()}", styles)

    headers = ["Date", "Statut", "Cote", "Cote suivante", "Message"]
    _abhl_v2_write_header_row(ws, 3, headers, styles)

    for index, row in enumerate(preview.get("rows") or [], start=4):
        inputs = row.get("inputs") or {}
        status = _abhl_v2_status_for_row(row)

        if status == "OK":
            message = "Ligne calculée."
        elif status == "INCOMPLET":
            message = "Ligne partiellement saisie."
        else:
            message = "Ligne non saisie."

        values = [
            _abhl_v2_to_date(row.get("date_bilan")),
            status,
            inputs.get("cote_7h_ngm"),
            inputs.get("cote_suivante_ngm"),
            message,
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col_idx, value=value)
            _abhl_v2_style_cell(cell, border=styles["border"], alignment=styles["center"])

            if col_idx == 1:
                cell.number_format = "dd/mm/yyyy"

    _abhl_v2_auto_width(ws)

    # On garde la feuille de contrôle disponible dans le fichier, mais cachée
    # pour ne pas perturber la structure officielle.
    ws.sheet_state = "hidden"


def generate_dynamic_bilan_workbook(db, output_path, preview: dict, config: dict | None = None):
    """
    Génère un fichier BILAN mensuel dynamique propre pour un nouveau barrage.

    Cette version est générique :
    - les colonnes de restitutions dépendent de preview["input_fields"]
    - elle fonctionne pour n'importe quel nouveau barrage
    - les anciens barrages restent sur leurs fichiers officiels
    """
    from openpyxl import Workbook

    wb = Workbook()

    default = wb.active
    wb.remove(default)

    styles = _abhl_v2_make_styles()

    _abhl_v2_write_guide_sheet(wb, preview, styles)
    _abhl_v2_write_evaporation_sheet(wb, preview, styles)
    _abhl_v2_write_apport_sheet(wb, preview, styles)
    _abhl_v2_write_bareme_sheet(wb, db, preview, styles)
    _abhl_v2_write_controle_sheet(wb, preview, styles)

    # Ordre officiel visible : Guide, EVAPORATION, APPORT, BAREME.
    wb.active = 0

    wb.save(output_path)


# === ABHL DYNAMIC BILAN V2 PATCH END ===
