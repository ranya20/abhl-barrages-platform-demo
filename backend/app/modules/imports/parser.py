from __future__ import annotations

import csv
import io
import re
import unicodedata
from calendar import monthrange
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.modules.annonce.mappings import SHEET_CONFIGS
from app.modules.bilan.mappings import BILAN_CONFIGS, get_bilan_config
from app.modules.imports.constants import (
    MODULE_ANNONCE,
    MODULE_BILAN,
    RESTITUTION_CODE_ALIASES,
)


HEADER_ALIASES = {
    "barrage": "barrage_code",
    "code_barrage": "barrage_code",
    "barrage_code": "barrage_code",
    "code": "barrage_code",
    "date": "date_bilan",
    "jour": "date_bilan",
    "date_bilan": "date_bilan",
    "date_situation": "date_situation",
    "cote": "cote_7h_ngm",
    "cote_interval": "cote_interval_ngm",
    "cote_interval_ngm": "cote_interval_ngm",
    "cote_j": "cote_interval_ngm",
    "cote_7h": "cote_7h_ngm",
    "cote_a_7h": "cote_7h_ngm",
    "cote_7h_ngm": "cote_7h_ngm",
    "cote_suivante": "cote_suivante_ngm",
    "cote_j1": "cote_suivante_ngm",
    "cote_suivante_ngm": "cote_suivante_ngm",
    "hauteur_bac": "hauteur_bac_mm",
    "hauteur_bac_mm": "hauteur_bac_mm",
    "h_bac": "hauteur_bac_mm",
    "pluie": "pluie_mm",
    "pluie_mm": "pluie_mm",
    "observation": "observation",
    "observations": "observation",
}


def normalize_text(value: Any) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text


def normalize_code(value: Any) -> str:
    code = normalize_text(value).upper()
    return RESTITUTION_CODE_ALIASES.get(code, code)


def normalize_header(value: Any) -> str:
    normalized = normalize_text(value)
    if normalized in HEADER_ALIASES:
        return HEADER_ALIASES[normalized]

    for prefix in ("restitution_", "rest_", "sortie_", "lacher_"):
        if normalized.startswith(prefix):
            return f"rest__{normalize_code(normalized[len(prefix):])}"

    return normalized


def as_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    elif "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def as_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _row_payload(
    *,
    row_number: int,
    sheet_name: str,
    barrage_code: str | None,
    date_bilan: date | None,
    data: dict[str, Any],
    raw: dict[str, Any] | None = None,
) -> dict:
    return {
        "row_number": row_number,
        "sheet_name": sheet_name,
        "barrage_code": normalize_code(barrage_code) if barrage_code else None,
        "date_bilan": date_bilan,
        "data": data,
        "raw": raw or data,
        "errors": [],
        "warnings": [],
    }


def _detect_header_row(values: list[list[Any]]) -> int | None:
    known = set(HEADER_ALIASES) | set(HEADER_ALIASES.values())
    for index, row in enumerate(values[:25]):
        normalized = [normalize_header(cell) for cell in row if cell not in (None, "")]
        score = sum(1 for cell in normalized if cell in known or cell.startswith("rest__"))
        if score >= 2:
            return index
    return None


def _generic_rows_from_matrix(
    values: list[list[Any]],
    *,
    sheet_name: str,
    default_barrage_code: str | None,
    default_date_situation: date | None,
) -> list[dict]:
    header_index = _detect_header_row(values)
    if header_index is None:
        return []

    headers = [normalize_header(value) for value in values[header_index]]
    result: list[dict] = []

    for offset, row in enumerate(values[header_index + 1 :], start=header_index + 2):
        if not any(value not in (None, "") for value in row):
            continue
        raw = {
            headers[index]: row[index]
            for index in range(min(len(headers), len(row)))
            if headers[index]
        }
        barrage_code = raw.get("barrage_code") or default_barrage_code
        row_date = as_date(raw.get("date_bilan"))
        date_situation = as_date(raw.get("date_situation")) or default_date_situation

        data: dict[str, Any] = {
            "date_situation": date_situation,
            "cote_interval_ngm": as_float(raw.get("cote_interval_ngm")),
            "cote_7h_ngm": as_float(raw.get("cote_7h_ngm")),
            "cote_suivante_ngm": as_float(raw.get("cote_suivante_ngm")),
            "hauteur_bac_mm": as_float(raw.get("hauteur_bac_mm")),
            "pluie_mm": as_float(raw.get("pluie_mm")),
            "observation": str(raw.get("observation") or "").strip() or None,
            "restitutions": {},
        }

        for key, value in raw.items():
            if key.startswith("rest__"):
                code = normalize_code(key.split("__", 1)[1])
                numeric = as_float(value)
                if numeric is not None:
                    data["restitutions"][code] = numeric

        item = _row_payload(
            row_number=offset,
            sheet_name=sheet_name,
            barrage_code=barrage_code,
            date_bilan=row_date,
            data=data,
            raw={key: _json_safe(value) for key, value in raw.items()},
        )
        result.append(item)

    return result


def _parse_csv(content: bytes, context: dict) -> list[dict]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"

    matrix = list(csv.reader(io.StringIO(text), dialect))
    return _generic_rows_from_matrix(
        matrix,
        sheet_name="CSV",
        default_barrage_code=context.get("barrage_code"),
        default_date_situation=as_date(context.get("date_situation")),
    )


def _parse_generic_workbook(workbook, context: dict) -> list[dict]:
    result: list[dict] = []
    for worksheet in workbook.worksheets:
        values = [list(row) for row in worksheet.iter_rows(values_only=True)]
        rows = _generic_rows_from_matrix(
            values,
            sheet_name=worksheet.title,
            default_barrage_code=context.get("barrage_code"),
            default_date_situation=as_date(context.get("date_situation")),
        )
        if rows:
            result.extend(rows)
    return result


def _workbook_period(workbook) -> tuple[int | None, int | None]:
    """Retourne (année, mois) depuis la page d'accueil du classeur officiel."""
    if "PAGE D'ACCEUIL" not in workbook.sheetnames:
        return None, None
    worksheet = workbook["PAGE D'ACCEUIL"]
    month_value = as_float(worksheet["F8"].value)
    year_value = as_float(worksheet["F9"].value)
    if month_value is None or year_value is None:
        return None, None
    month = int(month_value)
    year = int(year_value)
    if not (1 <= month <= 12 and 1900 <= year <= 2200):
        return None, None
    return year, month


def _resolve_dynamic_barrage_code(sheet_name: str, context: dict) -> str | None:
    aliases = {
        normalize_text(key): normalize_code(value)
        for key, value in dict(context.get("barrage_aliases") or {}).items()
    }
    normalized_sheet = normalize_text(sheet_name)
    if normalized_sheet in aliases:
        return aliases[normalized_sheet]

    candidate = normalize_code(sheet_name)
    known_codes = {normalize_code(value) for value in context.get("barrage_codes") or []}
    if candidate in known_codes:
        return candidate
    return None


def _dynamic_sheet_columns(worksheet, context: dict) -> tuple[int, dict[str, int], dict[int, str]] | None:
    """Détecte les colonnes d'une feuille de barrage non codée en dur."""
    restitution_aliases = {
        normalize_text(key): normalize_code(value)
        for key, value in dict(context.get("restitution_aliases") or {}).items()
    }
    restitution_codes = {normalize_code(value) for value in context.get("restitution_codes") or []}

    for header_row in range(1, min(worksheet.max_row, 18) + 1):
        field_columns: dict[str, int] = {}

        # Les champs structurants doivent être présents sur la ligne d'en-tête elle-même.
        for column in range(1, worksheet.max_column + 1):
            label = normalize_text(worksheet.cell(row=header_row, column=column).value)
            if not label:
                continue
            if "date" in label and "date" not in field_columns:
                field_columns["date"] = column
            elif "cote" in label and "cote" not in field_columns:
                field_columns["cote"] = column
            elif "hauteur" in label and "bac" in label and "hauteur_bac" not in field_columns:
                field_columns["hauteur_bac"] = column
            elif "pluie" in label and "pluie" not in field_columns:
                field_columns["pluie"] = column
            elif "observation" in label and "observation" not in field_columns:
                field_columns["observation"] = column

        if "date" not in field_columns or "cote" not in field_columns:
            continue

        rest_columns: dict[int, str] = {}
        for column in range(1, worksheet.max_column + 1):
            if column in field_columns.values():
                continue
            labels = []
            for row in range(header_row, min(header_row + 2, worksheet.max_row) + 1):
                value = worksheet.cell(row=row, column=column).value
                if value not in (None, ""):
                    labels.append(str(value))
            if not labels:
                continue

            rest_code = None
            for value in reversed(labels):
                label = normalize_text(value)
                if label in restitution_aliases:
                    rest_code = restitution_aliases[label]
                    break
                candidate = normalize_code(label)
                if candidate in restitution_codes:
                    rest_code = candidate
                    break
            if rest_code:
                rest_columns[column] = rest_code

        return header_row, field_columns, rest_columns

    return None

def _parse_dynamic_annonce_sheet(workbook, worksheet, context: dict) -> list[dict]:
    code = _resolve_dynamic_barrage_code(worksheet.title, context)
    if not code:
        return []

    detected = _dynamic_sheet_columns(worksheet, context)
    if not detected:
        return []
    header_row, fields, rest_columns = detected
    year, month = _workbook_period(workbook)

    first_data_row = None
    for row in range(header_row + 1, min(worksheet.max_row, header_row + 8) + 1):
        raw_date = as_date(worksheet.cell(row=row, column=fields["date"]).value)
        cote = as_float(worksheet.cell(row=row, column=fields["cote"]).value)
        if raw_date or cote is not None:
            first_data_row = row
            break
    if first_data_row is None:
        return []

    max_days = monthrange(year, month)[1] if year and month else 31
    result: list[dict] = []
    for offset in range(max_days):
        interval_row = first_data_row + offset
        current_row = interval_row + 1
        if current_row > worksheet.max_row:
            break

        interval_date = as_date(worksheet.cell(interval_row, fields["date"]).value)
        if interval_date is None and year and month:
            interval_date = date(year, month, offset + 1)
        if interval_date is None:
            continue

        current_date = as_date(worksheet.cell(current_row, fields["date"]).value) or (interval_date + timedelta(days=1))
        data = {
            "date_situation": current_date,
            "cote_interval_ngm": as_float(worksheet.cell(interval_row, fields["cote"]).value),
            "cote_7h_ngm": as_float(worksheet.cell(current_row, fields["cote"]).value),
            "cote_suivante_ngm": None,
            "hauteur_bac_mm": as_float(worksheet.cell(interval_row, fields.get("hauteur_bac", 0)).value)
            if fields.get("hauteur_bac") else None,
            "pluie_mm": as_float(worksheet.cell(interval_row, fields.get("pluie", 0)).value)
            if fields.get("pluie") else None,
            "observation": str(worksheet.cell(interval_row, fields.get("observation", 0)).value or "").strip() or None
            if fields.get("observation") else None,
            "restitutions": {},
        }
        for column, rest_code in rest_columns.items():
            value = as_float(worksheet.cell(interval_row, column).value)
            if value is not None:
                data["restitutions"][rest_code] = value

        if not any(
            value is not None
            for value in (data["hauteur_bac_mm"], data["pluie_mm"])
        ) and not data["restitutions"]:
            # Une ligne qui contient uniquement la cote sert de borne au jour précédent.
            continue

        result.append(
            _row_payload(
                row_number=interval_row,
                sheet_name=worksheet.title,
                barrage_code=code,
                date_bilan=interval_date,
                data=data,
                raw={
                    "interval_row": interval_row,
                    "current_row": current_row,
                    "date_interval": interval_date.isoformat(),
                    "date_situation": current_date.isoformat(),
                    "dynamic_sheet": True,
                },
            )
        )
    return result


def _parse_annonce_official(workbook, context: dict) -> list[dict]:
    """Analyse toutes les journées et toutes les feuilles d'un fichier Annonce."""
    year, month = _workbook_period(workbook)
    days_in_month = monthrange(year, month)[1] if year and month else 31
    result: list[dict] = []
    handled_sheets: set[str] = set()

    for code, config in SHEET_CONFIGS.items():
        sheet_name = config["sheet_name"]
        if sheet_name not in workbook.sheetnames:
            continue
        handled_sheets.add(sheet_name)
        worksheet = workbook[sheet_name]
        start_row = int(config["data_start_row"])

        for offset in range(days_in_month):
            interval_row = start_row + offset
            current_row = interval_row + 1
            if current_row > worksheet.max_row:
                break

            interval_date = as_date(worksheet[f"A{interval_row}"].value)
            if interval_date is None and year and month:
                interval_date = date(year, month, offset + 1)
            if interval_date is None:
                continue
            current_date = as_date(worksheet[f"A{current_row}"].value) or (interval_date + timedelta(days=1))

            data = {
                "date_situation": current_date,
                "cote_interval_ngm": as_float(worksheet[f"B{interval_row}"].value),
                "cote_7h_ngm": as_float(worksheet[f"B{current_row}"].value),
                "cote_suivante_ngm": None,
                "hauteur_bac_mm": as_float(worksheet[f"D{interval_row}"].value),
                "pluie_mm": as_float(worksheet[f"E{interval_row}"].value),
                "observation": None,
                "restitutions": {},
            }

            for type_code, column in config.get("restitutions", {}).items():
                value = as_float(worksheet[f"{column}{interval_row}"].value)
                if value is not None:
                    data["restitutions"][normalize_code(type_code)] = value

            transfer_column = config.get("specials", {}).get("TRANSFERT_DAR_KHROFA")
            if transfer_column:
                value = as_float(worksheet[f"{transfer_column}{interval_row}"].value)
                if value is not None:
                    data["restitutions"]["TRANSFERT_DAR_KHROFA"] = value

            if not any(
                value is not None
                for value in (data["hauteur_bac_mm"], data["pluie_mm"])
            ) and not data["restitutions"]:
                # Une ligne qui contient uniquement la cote sert de borne au jour précédent.
                continue

            result.append(
                _row_payload(
                    row_number=interval_row,
                    sheet_name=sheet_name,
                    barrage_code=code,
                    date_bilan=interval_date,
                    data=data,
                    raw={
                        "interval_row": interval_row,
                        "current_row": current_row,
                        "date_interval": interval_date.isoformat(),
                        "date_situation": current_date.isoformat(),
                    },
                )
            )

    ignored = {"PAGE D'ACCEUIL", "ANNONCE DE CRUES", "Baréms des barrages"}
    for worksheet in workbook.worksheets:
        if worksheet.title in handled_sheets or worksheet.title in ignored:
            continue
        result.extend(_parse_dynamic_annonce_sheet(workbook, worksheet, context))

    result.sort(key=lambda item: (item.get("date_bilan") or date.min, item.get("barrage_code") or "", item.get("row_number") or 0))
    return result

def _parse_bilan_official(workbook, context: dict) -> list[dict]:
    barrage_code = normalize_code(context.get("barrage_code"))
    if not barrage_code:
        raise ValueError("Le barrage sélectionné est obligatoire pour importer un BILAN officiel.")

    config = get_bilan_config(barrage_code)
    if not config:
        return []
    if "APPORT" not in workbook.sheetnames or "EVAPORATION" not in workbook.sheetnames:
        return []

    apport_ws = workbook["APPORT"]
    evaporation_ws = workbook["EVAPORATION"]
    apport_map = config["apport"]
    evaporation_map = config["evaporation"]

    evaporation_by_date: dict[date, dict] = {}
    for offset in range(int(evaporation_map.get("row_count", 32))):
        row = int(evaporation_map["start_row"]) + offset
        day = as_date(evaporation_ws[f"{evaporation_map['columns']['date']}{row}"].value)
        if not day:
            continue
        evaporation_by_date[day] = {
            "cote_7h_ngm": as_float(evaporation_ws[f"{evaporation_map['columns']['cote']}{row}"].value),
            "hauteur_bac_mm": as_float(evaporation_ws[f"{evaporation_map['columns']['hauteur_bac']}{row}"].value),
            "pluie_mm": as_float(evaporation_ws[f"{evaporation_map['columns']['pluie']}{row}"].value),
        }

    result: list[dict] = []
    for offset in range(int(apport_map.get("row_count", 32))):
        row = int(apport_map["start_row"]) + offset
        day = as_date(apport_ws[f"{apport_map['columns']['date']}{row}"].value)
        if not day:
            continue
        evap = evaporation_by_date.get(day, {})
        data = {
            "date_situation": None,
            "cote_7h_ngm": evap.get("cote_7h_ngm")
            if evap.get("cote_7h_ngm") is not None
            else as_float(apport_ws[f"{apport_map['columns']['cote']}{row}"].value),
            "cote_suivante_ngm": None,
            "hauteur_bac_mm": evap.get("hauteur_bac_mm"),
            "pluie_mm": evap.get("pluie_mm")
            if evap.get("pluie_mm") is not None
            else as_float(apport_ws[f"{apport_map['columns']['pluie']}{row}"].value),
            "observation": None,
            "restitutions": {},
        }

        for field in config.get("input_fields", []):
            value = as_float(apport_ws[f"{field['column']}{row}"].value)
            if value is not None:
                data["restitutions"][normalize_code(field["code"])] = value

        if not any(
            value is not None
            for value in (
                data["cote_7h_ngm"],
                data["hauteur_bac_mm"],
                data["pluie_mm"],
            )
        ) and not data["restitutions"]:
            continue

        result.append(
            _row_payload(
                row_number=row,
                sheet_name="APPORT",
                barrage_code=barrage_code,
                date_bilan=day,
                data=data,
                raw={"apport_row": row, "date_bilan": day.isoformat()},
            )
        )

    result.sort(key=lambda item: item["date_bilan"] or date.min)
    for index, item in enumerate(result):
        if index + 1 < len(result):
            item["data"]["cote_suivante_ngm"] = result[index + 1]["data"].get("cote_7h_ngm")
    return result


def _json_safe(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def parse_import_file(
    *,
    module_code: str,
    filename: str,
    content: bytes,
    context: dict,
) -> tuple[list[dict], str]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        rows = _parse_csv(content, context)
        return rows, "CSV_TABULAR"

    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)

    if module_code == MODULE_ANNONCE:
        official_sheet_names = {cfg["sheet_name"] for cfg in SHEET_CONFIGS.values()}
        has_dynamic_barrage_sheet = any(
            _resolve_dynamic_barrage_code(sheet_name, context)
            for sheet_name in workbook.sheetnames
        )
        if (
            official_sheet_names.intersection(workbook.sheetnames)
            or "PAGE D'ACCEUIL" in workbook.sheetnames
            or has_dynamic_barrage_sheet
        ):
            rows = _parse_annonce_official(workbook, context)
            if rows:
                return rows, "ANNONCE_OFFICIEL_COMPLET"

    if module_code == MODULE_BILAN and {"APPORT", "EVAPORATION"}.issubset(workbook.sheetnames):
        rows = _parse_bilan_official(workbook, context)
        if rows:
            return rows, "BILAN_OFFICIEL"

    rows = _parse_generic_workbook(workbook, context)
    return rows, "XLSX_TABULAR"
