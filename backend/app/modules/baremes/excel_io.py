from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_POINTS = 250_000
MIN_POINTS = 2

EXPECTED_HEADERS = ("Cote NGM", "Volume Mm3", "Surface km2")


def _norm(value) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = (
        text.replace("³", "3")
        .replace("²", "2")
        .replace("(", " ")
        .replace(")", " ")
        .replace("_", " ")
        .replace("-", " ")
    )
    return " ".join(text.lower().split())


HEADER_ALIASES = {
    "cote": {
        _norm("Cote NGM"),
        _norm("Côte NGM"),
        _norm("cote_ngm"),
        _norm("cote"),
    },
    "volume": {
        _norm("Volume Mm3"),
        _norm("Volume Mm³"),
        _norm("volume_mm3"),
        _norm("volume"),
    },
    "surface": {
        _norm("Surface km2"),
        _norm("Surface km²"),
        _norm("surface_km2"),
        _norm("surface"),
    },
}


def _decimal(value, *, row: int, field: str) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Ligne {row}: {field} est obligatoire.")
    try:
        if isinstance(value, str):
            value = value.strip().replace(" ", "").replace(",", ".")
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError(f"Ligne {row}: {field} doit être numérique.") from exc
    if not result.is_finite():
        raise ValueError(f"Ligne {row}: {field} n'est pas une valeur numérique finie.")
    return result


def _find_header(ws) -> tuple[int, dict[str, int]]:
    for row in range(1, min(ws.max_row, 30) + 1):
        found: dict[str, int] = {}
        for col in range(1, min(ws.max_column, 30) + 1):
            token = _norm(ws.cell(row=row, column=col).value)
            if not token:
                continue
            for key, aliases in HEADER_ALIASES.items():
                if token in aliases and key not in found:
                    found[key] = col
        if len(found) == 3:
            return row, found
    raise ValueError(
        "Structure Excel non reconnue. Utilisez le modèle ABHL avec les colonnes "
        "'Cote NGM', 'Volume Mm3' et 'Surface km2'."
    )


def validate_points(points: list[dict]) -> dict:
    if len(points) < MIN_POINTS:
        raise ValueError(f"Le barème doit contenir au moins {MIN_POINTS} points.")
    if len(points) > MAX_POINTS:
        raise ValueError(f"Le barème dépasse la limite de sécurité de {MAX_POINTS} points.")

    sorted_points = sorted(points, key=lambda p: p["cote_ngm"])
    errors: list[str] = []
    warnings: list[str] = []
    seen: dict[Decimal, dict] = {}
    previous = None

    for index, point in enumerate(sorted_points, start=1):
        cote = point["cote_ngm"]
        volume = point["volume_mm3"]
        surface = point["surface_km2"]

        if volume < 0:
            errors.append(f"Point {index}: volume négatif à la cote {cote}.")
        if surface < 0:
            errors.append(f"Point {index}: surface négative à la cote {cote}.")
        if cote in seen:
            errors.append(f"Cote dupliquée: {cote}.")
        seen[cote] = point

        if previous is not None:
            if volume < previous["volume_mm3"]:
                errors.append(
                    f"Volume décroissant entre les cotes "
                    f"{previous['cote_ngm']} et {cote}."
                )
            if surface < previous["surface_km2"]:
                warnings.append(
                    f"Surface décroissante entre les cotes "
                    f"{previous['cote_ngm']} et {cote}; vérifier le fichier officiel."
                )
        previous = point

    if errors:
        raise ValueError(" ".join(errors[:20]))

    canonical = "\n".join(
        f"{p['cote_ngm']}|{p['volume_mm3']}|{p['surface_km2']}"
        for p in sorted_points
    ).encode("utf-8")

    return {
        "points": sorted_points,
        "warnings": warnings[:100],
        "points_sha256": sha256(canonical).hexdigest(),
        "min_cote": sorted_points[0]["cote_ngm"],
        "max_cote": sorted_points[-1]["cote_ngm"],
        "min_volume": min(p["volume_mm3"] for p in sorted_points),
        "max_volume": max(p["volume_mm3"] for p in sorted_points),
        "min_surface": min(p["surface_km2"] for p in sorted_points),
        "max_surface": max(p["surface_km2"] for p in sorted_points),
    }


def parse_bareme_excel(content: bytes, filename: str | None = None) -> dict:
    if not content:
        raise ValueError("Le fichier Excel est vide.")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError(
            f"Le fichier dépasse {MAX_IMPORT_BYTES // (1024 * 1024)} Mo."
        )

    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Impossible de lire le fichier .xlsx.") from exc

    try:
        if "BAREME" in wb.sheetnames:
            ws = wb["BAREME"]
        else:
            ws = wb[wb.sheetnames[0]]

        header_row, cols = _find_header(ws)
        points: list[dict] = []

        for row in range(header_row + 1, ws.max_row + 1):
            values = [
                ws.cell(row=row, column=cols["cote"]).value,
                ws.cell(row=row, column=cols["volume"]).value,
                ws.cell(row=row, column=cols["surface"]).value,
            ]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue

            cote = _decimal(values[0], row=row, field="Cote NGM")
            volume = _decimal(values[1], row=row, field="Volume Mm3")
            surface = _decimal(values[2], row=row, field="Surface km2")
            points.append(
                {
                    "cote_ngm": cote,
                    "volume_mm3": volume,
                    "surface_km2": surface,
                    "ligne_source": row,
                }
            )

        report = validate_points(points)
        report.update(
            {
                "filename": filename or "bareme.xlsx",
                "sheet_name": ws.title,
                "source_sha256": sha256(content).hexdigest(),
                "header_row": header_row,
                "points_count": len(report["points"]),
            }
        )
        return report
    finally:
        wb.close()


def build_blank_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BAREME"

    headers = list(EXPECTED_HEADERS)
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:C1"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20

    decimal_validation = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=False,
    )
    decimal_validation.error = "Saisissez une valeur numérique positive ou nulle."
    decimal_validation.errorTitle = "Valeur invalide"
    ws.add_data_validation(decimal_validation)
    decimal_validation.add("B2:C250001")

    instructions = wb.create_sheet("INSTRUCTIONS")
    rows = [
        ("MODELE BAREME ABHL", ""),
        ("Structure obligatoire", "Cote NGM | Volume Mm3 | Surface km2"),
        ("Règle 1", "Une ligne = un point de barème."),
        ("Règle 2", "Les trois valeurs sont obligatoires."),
        ("Règle 3", "Ne dupliquez pas une cote."),
        ("Règle 4", "Le volume et la surface doivent être positifs ou nuls."),
        ("Règle 5", "Conservez toutes les décimales du fichier officiel."),
        ("Règle 6", "Ne renommez pas les trois colonnes du modèle."),
        ("Important", "L'import crée d'abord un BROUILLON. Il ne devient applicable qu'après publication."),
    ]
    for row_idx, (left, right) in enumerate(rows, start=1):
        instructions.cell(row=row_idx, column=1, value=left).font = Font(
            bold=row_idx in (1, 2)
        )
        instructions.cell(row=row_idx, column=2, value=right)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 95
    instructions.sheet_view.showGridLines = False

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
